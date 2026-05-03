"""Tests for extractor.py — uses sample docs generated in-memory."""
from __future__ import annotations

import io
import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from auditcheckr.extractor import extract
from auditcheckr.models import AnnotationKind


def _make_simple_docx(paragraphs: list[str]) -> bytes:
    """Create an in-memory .docx with the given paragraph texts."""
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_extract_basic_paragraphs():
    raw = _make_simple_docx(["First paragraph.", "Second paragraph."])
    model = extract(raw, source_id="test.docx")

    assert len(model.blocks) == 2
    assert model.blocks[0].text == "First paragraph."
    assert model.blocks[1].text == "Second paragraph."
    assert model.source_id == "test.docx"


def test_extract_empty_paragraphs_skipped():
    raw = _make_simple_docx(["Hello", "", "World"])
    model = extract(raw, source_id="test.docx")
    texts = [b.text for b in model.blocks]
    assert "" not in texts
    assert "Hello" in texts
    assert "World" in texts


def test_extract_heading_block_type():
    doc = Document()
    doc.add_heading("Contract Title", level=1)
    doc.add_paragraph("Body text here.")
    buf = io.BytesIO()
    doc.save(buf)

    model = extract(buf.getvalue(), source_id="headings.docx")
    heading_block = next((b for b in model.blocks if "Title" in b.text), None)
    assert heading_block is not None
    assert heading_block.block_type == "heading"


def test_extract_no_annotations_by_default():
    raw = _make_simple_docx(["Just plain text."])
    model = extract(raw, source_id="plain.docx")
    assert all(len(b.annotations) == 0 for b in model.blocks)


def test_extract_table_cells():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Cell A1"
    table.cell(0, 1).text = "Cell B1"
    table.cell(1, 0).text = "Cell A2"
    table.cell(1, 1).text = "Cell B2"
    buf = io.BytesIO()
    doc.save(buf)

    model = extract(buf.getvalue(), source_id="table.docx")
    texts = [b.text for b in model.blocks]
    assert "Cell A1" in texts
    assert "Cell B2" in texts
    assert all(
        b.block_type == "table_cell"
        for b in model.blocks
        if b.block_type == "table_cell"
    )
